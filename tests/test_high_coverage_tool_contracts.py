"""High-signal branch coverage for external-engine built-in tools using fakes."""
from __future__ import annotations

import json
import builtins
from pathlib import Path
import shutil
import sys
from types import ModuleType, SimpleNamespace

from PIL import Image
import pikepdf
import pytest

from pdf_agent.core import ErrorCode, ToolError
from pdf_agent.tools._builtins import (
    add_blank_pages,
    auto_rotate,
    compress,
    compare,
    flatten,
    form_fill,
    linearize,
    nup,
    office_to_pdf,
    ocr,
    pages_to_zip,
    pdf_to_html,
    pdf_to_images,
    pdf_to_office,
    pdf_to_pdfa,
    pdf_to_text,
    pdf_to_word,
    redact,
    remove_blank_pages,
    repair,
    signature,
    signature_info,
    split,
    tile_pages,
    validate,
    watermark_image,
)
from pdf_agent.tools._builtins.add_blank_pages import AddBlankPagesTool
from pdf_agent.tools._builtins.add_page_numbers import _make_number_overlay
from pdf_agent.tools._builtins.auto_rotate import AutoRotateTool
from pdf_agent.tools._builtins.barcode import BarcodeTool
from pdf_agent.tools._builtins.booklet import BookletTool
from pdf_agent.tools._builtins.compare import CompareTool
from pdf_agent.tools._builtins.compress import CompressTool
from pdf_agent.tools._builtins.crop import CropTool
from pdf_agent.tools._builtins.decrypt import DecryptTool
from pdf_agent.tools._builtins.delete import DeleteTool
from pdf_agent.tools._builtins.encrypt import EncryptTool
from pdf_agent.tools._builtins.extract_attachments import ExtractAttachmentsTool
from pdf_agent.tools._builtins.extract_images import ExtractImagesTool
from pdf_agent.tools._builtins.extract import ExtractTool
from pdf_agent.tools._builtins.flatten import FlattenTool
from pdf_agent.tools._builtins.form_fill import FormFillTool
from pdf_agent.tools._builtins.header_footer import HeaderFooterTool
from pdf_agent.tools._builtins.images_to_pdf import ImagesToPdfTool
from pdf_agent.tools._builtins.linearize import LinearizeTool
from pdf_agent.tools._builtins.merge import MergeTool
from pdf_agent.tools._builtins.metadata_info import MetadataInfoTool
from pdf_agent.tools._builtins.nup import NUpTool
from pdf_agent.tools._builtins.office_to_pdf import OfficeToPdfTool
from pdf_agent.tools._builtins.page_border import PageBorderTool
from pdf_agent.tools._builtins.pages_to_zip import PagesToZipTool
from pdf_agent.tools._builtins.ocr import OcrTool
from pdf_agent.tools._builtins.pdf_to_html import PdfToHtmlTool
from pdf_agent.tools._builtins.pdf_to_images import PdfToImagesTool
from pdf_agent.tools._builtins.pdf_to_markdown import PdfToMarkdownTool
from pdf_agent.tools._builtins.pdf_to_office import PdfToExcelTool, PdfToPptTool
from pdf_agent.tools._builtins.pdf_to_pdfa import PdfATool
from pdf_agent.tools._builtins.pdf_to_text import PdfToTextTool
from pdf_agent.tools._builtins.pdf_to_word import PdfToWordTool
from pdf_agent.tools._builtins.qr_code import QrCodeTool
from pdf_agent.tools._builtins.redact import RedactTool
from pdf_agent.tools._builtins.remove_blank_pages import RemoveBlankPagesTool
from pdf_agent.tools._builtins.remove_metadata import RemoveMetadataTool
from pdf_agent.tools._builtins.reorder import ReorderTool
from pdf_agent.tools._builtins.repair import RepairTool
from pdf_agent.tools._builtins.resize import ResizeTool
from pdf_agent.tools._builtins.signature import SignatureTool
from pdf_agent.tools._builtins.signature_info import SignatureInfoTool
from pdf_agent.tools._builtins.split import SplitTool
from pdf_agent.tools._builtins.stamp import StampTool
from pdf_agent.tools._builtins.tile_pages import TilePagesTool
from pdf_agent.tools._builtins.validate import ValidateTool
from pdf_agent.tools._builtins.watermark_image import WatermarkImageTool
from pdf_agent.tools._builtins.watermark_text import WatermarkTextTool
from pdf_agent.tools import libreoffice

pytestmark = pytest.mark.coverage_edges


def test_auto_rotate_detection_helpers_cover_engines_and_fallbacks(
    tmp_path: Path,
    sample_images: list[Path],
    monkeypatch: pytest.MonkeyPatch,
):
    assert auto_rotate._is_low_text_osd_error(ToolError(ErrorCode.ENGINE_EXEC_FAILED, "Too few characters")) is True
    assert auto_rotate._is_low_text_osd_error(ToolError(ErrorCode.INVALID_PARAMS, "Too few characters")) is False
    assert auto_rotate._nearest_right_angle(359) == 0
    assert auto_rotate._nearest_right_angle(91) == 90

    class Char:
        def __init__(self, text: str) -> None:
            self.text = text

        def get_text(self) -> str:
            return self.text

    class Container(list):
        pass

    nested = Container([Char("a"), Container([Char("b")])])
    assert [char.get_text() for char in auto_rotate._iter_pdf_chars(nested, Char, Container)] == ["a", "b"]

    monkeypatch.setattr(auto_rotate.shutil, "which", lambda name: None)
    assert auto_rotate._detect_rotation(sample_images[0]) == (0, None)
    assert auto_rotate._ocr_rotation_score(sample_images[0], 0) is None
    assert auto_rotate._render_page_png(Path("missing.pdf"), 0, tmp_path) is None

    def fake_which(name: str) -> str | None:
        return f"/usr/bin/{name}"

    monkeypatch.setattr(auto_rotate.shutil, "which", fake_which)
    monkeypatch.setattr(
        auto_rotate,
        "run_command",
        lambda *_args, **_kwargs: SimpleNamespace(
            returncode=0,
            stdout=b"Rotate: 90\nOrientation confidence: 12.5\n",
            stderr=b"",
        ),
    )
    assert auto_rotate._detect_rotation(sample_images[0]) == (90, 12.5)

    monkeypatch.setattr(
        auto_rotate,
        "run_command",
        lambda *_args, **_kwargs: SimpleNamespace(returncode=1, stdout=b"", stderr=b"Too few characters"),
    )
    with pytest.raises(ToolError) as osd_error:
        auto_rotate._detect_rotation(sample_images[0])
    assert osd_error.value.code == ErrorCode.ENGINE_EXEC_FAILED

    tsv = "level\tconf\ttext\n1\t91\tHelloWorld123\n2\t-1\tignored\n3\tbad\tignored\n"
    monkeypatch.setattr(
        auto_rotate,
        "run_command",
        lambda *_args, **_kwargs: SimpleNamespace(returncode=0, stdout=tsv.encode(), stderr=b""),
    )
    assert auto_rotate._ocr_rotation_score(sample_images[0], 90) == (91.0, 13)

    monkeypatch.setattr(auto_rotate, "_ocr_rotation_score", lambda _image, angle: (float(angle), 20) if angle in {0, 180} else None)
    assert auto_rotate._detect_rotation_with_ocr_fallback(sample_images[0]) == (180, 180.0)
    monkeypatch.setattr(auto_rotate, "_ocr_rotation_score", lambda _image, angle: None)
    assert auto_rotate._detect_rotation_with_ocr_fallback(sample_images[0]) == (0, None)


def test_auto_rotate_render_and_run_contracts(
    tmp_path: Path,
    sample_pdf: Path,
    monkeypatch: pytest.MonkeyPatch,
):
    monkeypatch.setattr(auto_rotate.shutil, "which", lambda name: f"/usr/bin/{name}")

    def render_command(cmd, **_kwargs):
        Path(str(cmd[-1]) + "-1.png").write_bytes(b"png")
        return SimpleNamespace(returncode=0, stdout=b"", stderr=b"")

    monkeypatch.setattr(auto_rotate, "run_command", render_command)
    assert auto_rotate._render_page_png(sample_pdf, 0, tmp_path).name == "p0-1.png"

    monkeypatch.setattr(
        auto_rotate,
        "run_command",
        lambda *_args, **_kwargs: SimpleNamespace(returncode=1, stdout=b"", stderr=b"render failed"),
    )
    with pytest.raises(ToolError) as render_failed:
        auto_rotate._render_page_png(sample_pdf, 0, tmp_path)
    assert render_failed.value.message == "render failed"

    monkeypatch.setattr(auto_rotate.shutil, "which", lambda name: None if name == "tesseract" else "/usr/bin/pdftoppm")
    with pytest.raises(ToolError) as no_tesseract:
        AutoRotateTool().run([sample_pdf], {}, tmp_path / "auto")
    assert no_tesseract.value.code == ErrorCode.ENGINE_NOT_INSTALLED

    monkeypatch.setattr(auto_rotate.shutil, "which", lambda name: "/usr/bin/tesseract" if name == "tesseract" else None)
    with pytest.raises(ToolError) as no_renderer:
        AutoRotateTool().run([sample_pdf], {}, tmp_path / "auto")
    assert no_renderer.value.code == ErrorCode.ENGINE_NOT_INSTALLED

    monkeypatch.setattr(auto_rotate.shutil, "which", lambda name: f"/usr/bin/{name}")
    monkeypatch.setattr(auto_rotate, "_render_page_png", lambda _pdf, index, td: td / f"p{index}.png")
    monkeypatch.setattr(auto_rotate, "_detect_rotation", lambda _image: (90, 50.0))
    updates: list[tuple[int, str]] = []
    result = AutoRotateTool().run(
        [sample_pdf],
        {"min_confidence": 10},
        tmp_path / "auto",
        reporter=lambda percent, message="": updates.append((percent, message)),
    )

    assert result.meta["total_pages"] == 5
    assert result.meta["rotations"][0]["angle"] == 90
    assert result.output_files[0].exists()
    assert updates[-1] == (100, "Done")


def test_redact_validation_text_box_and_rasterize_helpers(
    tmp_path: Path,
    sample_pdf: Path,
    monkeypatch: pytest.MonkeyPatch,
):
    tool = RedactTool()
    with pytest.raises(ToolError) as invalid_json:
        tool.validate({"regions_json": "{"})
    assert invalid_json.value.code == ErrorCode.INVALID_PARAMS
    assert redact._normalize_regions([
        {"page": "1", "x": "10", "y": "20", "width": "30", "height": "40"},
        {"bad": "ignored"},
    ]) == [{"page": 1, "x": 10.0, "y": 20.0, "width": 30.0, "height": 40.0}]

    monkeypatch.setattr(redact.shutil, "which", lambda name: None)
    with pytest.raises(ToolError) as no_pdftotext:
        redact._find_text_boxes(sample_pdf, "secret")
    assert no_pdftotext.value.code == ErrorCode.ENGINE_NOT_INSTALLED

    monkeypatch.setattr(redact.shutil, "which", lambda name: "/usr/bin/pdftotext")
    assert redact._find_text_boxes(sample_pdf, " , ") == []

    def fake_pdftotext(cmd, **_kwargs):
        xml_path = Path(cmd[-1])
        xml_path.write_text(
            """<doc><page height="100"><word xMin="10" yMin="20" xMax="40" yMax="35">secret</word></page></doc>""",
            encoding="utf-8",
        )
        return SimpleNamespace(returncode=0, stderr=b"")

    monkeypatch.setattr(redact, "run_command", fake_pdftotext)
    assert redact._find_text_boxes(sample_pdf, "secret") == [
        {"page": 1, "x": 10.0, "y": 65.0, "width": 30.0, "height": 15.0}
    ]

    def fake_gs(cmd, **_kwargs):
        output_args = [part for part in cmd if str(part).startswith("-sOutputFile=")]
        if output_args:
            output = Path(str(output_args[0]).split("=", 1)[1])
            if output.suffix == ".pdf":
                output.write_bytes(sample_pdf.read_bytes())
            else:
                output.write_bytes(b"png")
        return SimpleNamespace(returncode=0, stderr=b"")

    monkeypatch.setattr(redact, "run_command", fake_gs)
    out = tmp_path / "redacted.pdf"
    progress: list[int] = []
    redact._rasterize_pages(
        gs_bin="/usr/bin/gs",
        input_path=sample_pdf,
        output_path=out,
        redacted_pages={0},
        total_pages=1,
        workdir=tmp_path,
        reporter=lambda percent, message="": progress.append(percent),
    )
    assert out.exists()
    assert progress[-1] == 100


def test_redact_run_error_visual_full_and_rasterize_failure_paths(
    tmp_path: Path,
    sample_pdf: Path,
    monkeypatch: pytest.MonkeyPatch,
):
    tool = RedactTool()
    with pytest.raises(ToolError) as no_targets:
        tool.run([sample_pdf], {"regions_json": "[]", "text_query": ""}, tmp_path / "redact-empty")
    assert no_targets.value.code == ErrorCode.INVALID_PARAMS

    monkeypatch.setattr(redact.shutil, "which", lambda name: None)
    visual_dir = tmp_path / "redact-visual"
    visual_dir.mkdir()
    visual = tool.run(
        [sample_pdf],
        {"regions_json": '[{"page":1,"x":10,"y":10,"width":20,"height":20}]'},
        visual_dir,
    )
    assert visual.meta["content_removed"] is False
    assert visual.meta["redaction_mode"] == "visual_only"

    def successful_rasterize(*, output_path: Path, input_path: Path, **_kwargs):
        output_path.write_bytes(input_path.read_bytes())

    monkeypatch.setattr(redact.shutil, "which", lambda name: "/usr/bin/gs")
    monkeypatch.setattr(redact, "_rasterize_pages", successful_rasterize)
    full_dir = tmp_path / "redact-full"
    full_dir.mkdir()
    full = tool.run(
        [sample_pdf],
        {"regions_json": '[{"page":1,"x":10,"y":10,"width":20,"height":20}]'},
        full_dir,
    )
    assert full.meta["content_removed"] is True
    assert full.meta["redaction_mode"] == "full"

    def failing_rasterize(**_kwargs):
        raise RuntimeError("gs failed")

    monkeypatch.setattr(redact, "_rasterize_pages", failing_rasterize)
    fallback_dir = tmp_path / "redact-fallback"
    fallback_dir.mkdir()
    fallback = tool.run(
        [sample_pdf],
        {"regions_json": '[{"page":1,"x":10,"y":10,"width":20,"height":20}]'},
        fallback_dir,
    )
    assert "Ghostscript rasterization failed" in fallback.meta["warning"]


def test_office_to_pdf_libleoffice_success_and_python_fallbacks(
    tmp_path: Path,
    sample_docx: Path,
    sample_xlsx: Path,
    sample_pptx: Path,
    monkeypatch: pytest.MonkeyPatch,
):
    tool = OfficeToPdfTool()
    assert tool.validate({"ignored": "value"}) == {}
    assert tool.manifest().name == "office_to_pdf"

    monkeypatch.setattr(office_to_pdf.shutil, "which", lambda name: "/usr/bin/libreoffice")

    def fake_lo(_lo_bin, *, input_path: Path, outdir: Path, **_kwargs):
        outdir.mkdir(parents=True, exist_ok=True)
        (outdir / f"{input_path.stem}.pdf").write_bytes(b"%PDF-1.4\n%%EOF\n")
        return True, None

    monkeypatch.setattr(office_to_pdf, "run_libreoffice_conversion", fake_lo)
    progress: list[tuple[int, str]] = []
    lo_dir = tmp_path / "lo"
    lo_dir.mkdir()
    lo_result = tool.run(
        [sample_docx],
        {},
        lo_dir,
        reporter=lambda percent, message="": progress.append((percent, message)),
    )
    assert lo_result.meta["engine"] == "libreoffice"
    assert lo_result.meta["fallback_used"] is False
    assert progress[-1] == (100, "Done")

    monkeypatch.setattr(office_to_pdf.shutil, "which", lambda name: None)
    for source in (sample_docx, sample_xlsx, sample_pptx):
        fallback_dir = tmp_path / f"fallback-{source.suffix[1:]}"
        fallback_dir.mkdir()
        fallback = tool.run([source], {}, fallback_dir)
        assert fallback.meta["engine"] == "python-fallback"
        assert fallback.meta["fallback_used"] is True
        assert fallback.output_files[0].exists()

    unsupported = tmp_path / "notes.txt"
    unsupported.write_text("hello", encoding="utf-8")
    unsupported_dir = tmp_path / "fallback-unsupported"
    unsupported_dir.mkdir()
    with pytest.raises(ToolError) as no_fallback:
        tool.run([unsupported], {}, unsupported_dir)
    assert no_fallback.value.code == ErrorCode.OUTPUT_GENERATION_FAILED

    long_pdf = tmp_path / "long.pdf"
    office_to_pdf._render_text_lines_to_pdf(["x" * 200 for _ in range(80)], long_pdf)
    assert long_pdf.exists()


def test_ocr_modes_and_command_contracts(
    tmp_path: Path,
    sample_pdf: Path,
    monkeypatch: pytest.MonkeyPatch,
):
    tool = OcrTool()
    assert tool.validate({"skip_text": "false", "deskew": "yes", "output_mode": "json"}) == {
        "language": "eng",
        "skip_text": False,
        "deskew": True,
        "page_range": "all",
        "output_mode": "json",
    }
    monkeypatch.setattr(ocr.shutil, "which", lambda name: None)
    with pytest.raises(ToolError) as no_ocr:
        tool.run([sample_pdf], {}, tmp_path / "ocr-missing")
    assert no_ocr.value.code == ErrorCode.ENGINE_NOT_INSTALLED

    commands: list[list[str]] = []

    def fake_ocr_command(cmd, **_kwargs):
        commands.append([str(part) for part in cmd])
        output_path = Path(cmd[-1])
        output_path.parent.mkdir(parents=True, exist_ok=True)
        output_path.write_bytes(sample_pdf.read_bytes())
        sidecar = Path(cmd[cmd.index("--sidecar") + 1])
        sidecar.parent.mkdir(parents=True, exist_ok=True)
        sidecar.write_text("recognized text", encoding="utf-8")
        return SimpleNamespace(returncode=0, stderr=b"")

    monkeypatch.setattr(ocr.shutil, "which", lambda name: "/usr/bin/ocrmypdf")
    monkeypatch.setattr(ocr, "run_command", fake_ocr_command)
    pdf_result = tool.run(
        [sample_pdf],
        {"language": "chi_sim+eng", "skip_text": True, "deskew": True, "page_range": "1-2"},
        tmp_path / "ocr-pdf",
    )
    assert pdf_result.output_files[0].suffix == ".pdf"
    assert "--skip-text" in commands[-1]
    assert "--deskew" in commands[-1]
    assert "--pages" in commands[-1]

    txt_result = tool.run([sample_pdf], {"output_mode": "txt"}, tmp_path / "ocr-txt")
    assert txt_result.output_files[0].name == "ocr_output.txt"
    assert txt_result.output_files[0].read_text(encoding="utf-8") == "recognized text"

    json_result = tool.run([sample_pdf], {"output_mode": "json"}, tmp_path / "ocr-json")
    payload = json.loads(json_result.output_files[0].read_text(encoding="utf-8"))
    assert payload["text"] == "recognized text"


def test_form_fill_lists_validates_fills_and_flattens(
    tmp_path: Path,
    sample_pdf: Path,
    form_pdf: Path,
    monkeypatch: pytest.MonkeyPatch,
):
    tool = FormFillTool()
    assert tool.manifest().name == "form_fill"
    with pytest.raises(ToolError) as invalid_json:
        tool.validate({"field_values": "{"})
    assert invalid_json.value.code == ErrorCode.INVALID_PARAMS
    with pytest.raises(ToolError) as not_object:
        tool.validate({"field_values": "[]"})
    assert not_object.value.code == ErrorCode.INVALID_PARAMS

    list_result = tool.run([form_pdf], {}, tmp_path / "form-list")
    assert list_result.output_files == []
    assert "name" in list_result.meta["fields"]

    with pytest.raises(ToolError) as no_fields:
        tool.run([sample_pdf], {"field_values": '{"name":"Alice"}'}, tmp_path / "form-none")
    assert no_fields.value.code == ErrorCode.INVALID_PARAMS

    with pytest.raises(ToolError) as unknown:
        tool.run([form_pdf], {"field_values": '{"missing":"Alice"}'}, tmp_path / "form-missing")
    assert "Unknown form field" in unknown.value.message

    filled = tool.run([form_pdf], {"field_values": {"name": "Alice"}}, tmp_path / "form-fill")
    assert filled.meta["filled_fields"] == ["name"]
    assert filled.output_files[0].exists()

    def fake_flatten(source_path: Path, output_path: Path) -> None:
        output_path.write_bytes(source_path.read_bytes())

    monkeypatch.setattr(form_fill, "_flatten_pdf", fake_flatten)
    flattened = tool.run([form_pdf], {"field_values": {"name": "Bob"}, "flatten": True}, tmp_path / "form-flat")
    assert flattened.output_files[0].exists()


def test_images_to_pdf_validation_conversion_and_image_errors(
    tmp_path: Path,
    sample_images: list[Path],
):
    tool = ImagesToPdfTool()
    assert tool.manifest().name == "images_to_pdf"
    with pytest.raises(ToolError) as bad_size:
        tool.validate({"page_size": "A3"})
    assert bad_size.value.code == ErrorCode.INVALID_PARAMS
    with pytest.raises(ToolError) as no_images:
        tool.run([], {}, tmp_path / "images-empty")
    assert no_images.value.code == ErrorCode.INVALID_INPUT_FILE

    bad_image = tmp_path / "bad.png"
    bad_image.write_text("not an image", encoding="utf-8")
    with pytest.raises(ToolError) as bad_input:
        tool.run([bad_image], {}, tmp_path / "images-bad")
    assert bad_input.value.code == ErrorCode.INVALID_INPUT_FILE

    rgba = tmp_path / "rgba.png"
    Image.new("RGBA", (40, 40), (255, 0, 0, 128)).save(rgba)
    paletted = tmp_path / "paletted.gif"
    Image.new("P", (30, 30)).save(paletted)
    progress: list[int] = []
    images_dir = tmp_path / "images-pdf"
    images_dir.mkdir()
    result = tool.run(
        [rgba, paletted, sample_images[0]],
        {"page_size": "A4"},
        images_dir,
        reporter=lambda percent, message="": progress.append(percent),
    )
    assert result.output_files[0].exists()
    assert result.meta["image_count"] == 3
    assert progress[-1] == 100


def test_pdf_to_html_poppler_libreoffice_and_missing_engine_paths(
    tmp_path: Path,
    sample_pdf: Path,
    monkeypatch: pytest.MonkeyPatch,
):
    tool = PdfToHtmlTool()
    assert tool.validate({"single_page": False}) == {"single_page": False}
    monkeypatch.setattr(pdf_to_html.shutil, "which", lambda name: None)
    with pytest.raises(ToolError) as missing:
        tool.run([sample_pdf], {}, tmp_path / "html-missing")
    assert missing.value.code == ErrorCode.ENGINE_NOT_INSTALLED

    monkeypatch.setattr(pdf_to_html.shutil, "which", lambda name: "/usr/bin/pdftohtml" if name == "pdftohtml" else None)

    def fake_pdftohtml(cmd, **_kwargs):
        output_stem = Path(cmd[-1])
        output_stem.parent.mkdir(parents=True, exist_ok=True)
        (output_stem.parent / f"{output_stem.name}.html").write_text("<html>ok</html>", encoding="utf-8")
        return SimpleNamespace(returncode=0)

    monkeypatch.setattr(pdf_to_html, "run_command", fake_pdftohtml)
    updates: list[tuple[int, str]] = []
    poppler = tool.run(
        [sample_pdf],
        {"single_page": True},
        tmp_path / "html-poppler",
        reporter=lambda percent, message="": updates.append((percent, message)),
    )
    assert poppler.meta == {"engine": "pdftohtml", "files": 1}
    assert updates[-1] == (100, "Done")

    monkeypatch.setattr(pdf_to_html, "run_command", lambda *_args, **_kwargs: SimpleNamespace(returncode=0))
    with pytest.raises(ToolError) as no_output:
        tool.run([sample_pdf], {}, tmp_path / "html-no-output")
    assert no_output.value.code == ErrorCode.OUTPUT_GENERATION_FAILED

    monkeypatch.setattr(pdf_to_html.shutil, "which", lambda name: "/usr/bin/libreoffice" if name == "libreoffice" else None)
    def fake_lo_html(_lo, *, output_path, **_kwargs):
        output_path.parent.mkdir(parents=True, exist_ok=True)
        output_path.write_text("<html>lo</html>", encoding="utf-8")
        return True, None

    monkeypatch.setattr(pdf_to_html, "run_libreoffice_conversion_to_output", fake_lo_html)
    libre = tool.run([sample_pdf], {}, tmp_path / "html-lo")
    assert libre.meta == {"engine": "libreoffice"}

    monkeypatch.setattr(
        pdf_to_html,
        "run_libreoffice_conversion_to_output",
        lambda *_args, **_kwargs: (False, "lo failed"),
    )
    with pytest.raises(ToolError) as lo_failed:
        tool.run([sample_pdf], {}, tmp_path / "html-lo-fail")
    assert lo_failed.value.message == "lo failed"


def test_libreoffice_helpers_build_commands_and_normalize_outputs(
    tmp_path: Path,
    sample_docx: Path,
    monkeypatch: pytest.MonkeyPatch,
):
    cmd = libreoffice.build_libreoffice_command(
        "/usr/bin/libreoffice",
        convert_to="pdf",
        input_path=sample_docx,
        outdir=tmp_path / "out",
        profile_dir=tmp_path / "profile",
    )
    assert cmd[:4] == ["/usr/bin/libreoffice", f"-env:UserInstallation={(tmp_path / 'profile').resolve().as_uri()}", "--headless", "--convert-to"]

    calls: list[list[str]] = []
    monkeypatch.setattr(libreoffice, "run_command", lambda cmd, **_kwargs: calls.append(cmd))
    assert libreoffice.run_libreoffice_conversion(
        "/usr/bin/libreoffice",
        convert_to="pdf",
        input_path=sample_docx,
        outdir=tmp_path / "out",
        profile_dir=tmp_path / "profile",
        timeout=12,
    ) == (True, None)
    assert calls

    def fail_command(*_args, **_kwargs):
        raise ToolError(ErrorCode.ENGINE_EXEC_FAILED, "lo failed")

    monkeypatch.setattr(libreoffice, "run_command", fail_command)
    success, reason = libreoffice.run_libreoffice_conversion(
        "/usr/bin/libreoffice",
        convert_to="pdf",
        input_path=sample_docx,
        outdir=tmp_path / "out",
        profile_dir=tmp_path / "profile",
    )
    assert success is False
    assert "lo failed" in reason

    default_output = tmp_path / "out" / "sample.pdf"
    default_output.parent.mkdir(parents=True, exist_ok=True)
    default_output.write_bytes(b"%PDF-1.4\n%%EOF\n")
    monkeypatch.setattr(libreoffice, "run_libreoffice_conversion", lambda *args, **kwargs: (True, None))
    requested = tmp_path / "requested.pdf"
    assert libreoffice.run_libreoffice_conversion_to_output(
        "/usr/bin/libreoffice",
        convert_to="pdf",
        input_path=sample_docx,
        output_path=requested,
        outdir=default_output.parent,
        profile_dir=tmp_path / "profile",
    ) == (True, None)
    assert requested.exists()

    assert libreoffice.run_libreoffice_conversion_to_output(
        "/usr/bin/libreoffice",
        convert_to="pdf",
        input_path=sample_docx,
        output_path=requested,
        outdir=default_output.parent,
        profile_dir=tmp_path / "profile",
    ) == (True, None)

    missing = tmp_path / "missing.pdf"
    assert libreoffice.run_libreoffice_conversion_to_output(
        "/usr/bin/libreoffice",
        convert_to="pdf",
        input_path=sample_docx,
        output_path=missing,
        outdir=tmp_path / "empty",
        profile_dir=tmp_path / "profile",
    ) == (False, "LibreOffice did not produce a .pdf file")

    monkeypatch.setattr(libreoffice, "run_libreoffice_conversion", lambda *args, **kwargs: (False, "bad"))
    assert libreoffice.run_libreoffice_conversion_to_output(
        "/usr/bin/libreoffice",
        convert_to="pdf",
        input_path=sample_docx,
        output_path=missing,
        outdir=tmp_path / "empty",
        profile_dir=tmp_path / "profile",
    ) == (False, "bad")


def test_compare_render_and_run_branches(
    tmp_path: Path,
    sample_pdf: Path,
    monkeypatch: pytest.MonkeyPatch,
):
    import shutil
    import pdf_agent.external_commands as external_commands

    monkeypatch.setattr(shutil, "which", lambda name: None)
    assert compare._render_page_png(sample_pdf, 0) is None

    monkeypatch.setattr(shutil, "which", lambda name: "/usr/bin/pdftoppm")
    monkeypatch.setattr(
        external_commands,
        "run_command",
        lambda *_args, **_kwargs: SimpleNamespace(returncode=1, stderr=b"bad"),
    )
    assert compare._render_page_png(sample_pdf, 0) is None

    def fake_render_success(cmd, **_kwargs):
        out_stem = Path(cmd[-1])
        Image.new("RGB", (20, 20), "white").save(out_stem.parent / "page-1.png")
        return SimpleNamespace(returncode=0, stderr=b"")

    monkeypatch.setattr(external_commands, "run_command", fake_render_success)
    assert compare._render_page_png(sample_pdf, 0).size == (20, 20)

    workdir = tmp_path / "compare"
    workdir.mkdir()
    monkeypatch.setattr(compare, "_render_page_png", lambda *_args, **_kwargs: None)
    with pytest.raises(ToolError) as no_pages:
        CompareTool().run([sample_pdf, sample_pdf], {}, workdir)
    assert no_pages.value.code == ErrorCode.OUTPUT_GENERATION_FAILED

    def fake_page_image(_path: Path, page_idx: int, dpi: int = 72):
        if page_idx == 0:
            return Image.new("RGB", (20, 20), "white")
        return Image.new("RGB", (20, 20), "black")

    monkeypatch.setattr(compare, "_render_page_png", fake_page_image)
    monkeypatch.setattr(compare, "_extract_page_text", lambda _page: "text")
    updates: list[tuple[int, str]] = []
    ok_dir = tmp_path / "compare-ok"
    ok_dir.mkdir()
    result = CompareTool().run(
        [sample_pdf, sample_pdf],
        {"highlight_color": "blue", "sensitivity": 1},
        ok_dir,
        reporter=lambda percent, message="": updates.append((percent, message)),
    )
    assert result.meta["pages_compared"] == 5
    assert result.output_files[0].exists()
    assert result.output_files[1].name == "diff_text.json"
    assert updates[-1] == (100, "Done")


def test_pdf_to_images_validation_all_pages_ranges_and_output_collection(
    tmp_path: Path,
    sample_pdf: Path,
    monkeypatch: pytest.MonkeyPatch,
):
    tool = PdfToImagesTool()
    with pytest.raises(ToolError) as invalid:
        tool.validate({"format": "gif"})
    assert invalid.value.code == ErrorCode.INVALID_PARAMS

    monkeypatch.setattr(pdf_to_images.shutil, "which", lambda name: None)
    with pytest.raises(ToolError) as missing:
        tool.run([sample_pdf], {}, tmp_path / "images-missing")
    assert missing.value.code == ErrorCode.ENGINE_NOT_INSTALLED

    def fake_convert_command(cmd, **_kwargs):
        workdir = Path(cmd[-1]).parent
        ext = "jpg" if "-jpeg" in cmd else "png"
        (workdir / f"page-1.{ext}").write_bytes(b"image")
        return SimpleNamespace(returncode=0)

    monkeypatch.setattr(pdf_to_images.shutil, "which", lambda name: "/usr/bin/pdftoppm")
    monkeypatch.setattr(pdf_to_images, "run_command", fake_convert_command)
    all_dir = tmp_path / "pdf-images-all"
    all_dir.mkdir()
    all_pages = tool.run([sample_pdf], {"format": "jpeg", "dpi": 72}, all_dir)
    assert all_pages.meta["format"] == "jpeg"
    assert all_pages.output_files[0].suffix == ".jpg"

    range_dir = tmp_path / "pdf-images-range"
    range_dir.mkdir()
    subset = tool.run([sample_pdf], {"format": "png", "page_range": "1-2"}, range_dir)
    assert subset.meta["page_count"] == 2

    empty_dir = tmp_path / "pdf-images-empty"
    empty_dir.mkdir()
    monkeypatch.setattr(pdf_to_images, "run_command", lambda *_args, **_kwargs: SimpleNamespace(returncode=0))
    with pytest.raises(ToolError) as no_output:
        PdfToImagesTool._convert_pdf("/usr/bin/pdftoppm", sample_pdf, empty_dir, "png", 150)
    assert no_output.value.code == ErrorCode.OUTPUT_GENERATION_FAILED


def test_pdf_to_office_libreoffice_and_fallback_paths(
    tmp_path: Path,
    sample_pdf: Path,
    monkeypatch: pytest.MonkeyPatch,
):
    excel = PdfToExcelTool()
    ppt = PdfToPptTool()
    assert excel.validate({}) == {}
    assert ppt.validate({}) == {}

    monkeypatch.setattr(pdf_to_office.shutil, "which", lambda name: None)
    excel_dir = tmp_path / "excel-fallback"
    excel_dir.mkdir()
    excel_result = excel.run([sample_pdf], {}, excel_dir)
    assert excel_result.meta["engine"] == "openpyxl"
    assert excel_result.output_files[0].exists()

    ppt_dir = tmp_path / "ppt-fallback"
    ppt_dir.mkdir()
    ppt_result = ppt.run([sample_pdf], {}, ppt_dir)
    assert ppt_result.meta["engine"] == "python-pptx"
    assert ppt_result.output_files[0].exists()

    monkeypatch.setattr(pdf_to_office.shutil, "which", lambda name: "/usr/bin/libreoffice")

    def fake_success(_lo, *, output_path: Path, **_kwargs):
        output_path.write_bytes(b"converted")
        return True, None

    monkeypatch.setattr(pdf_to_office, "run_libreoffice_conversion_to_output", fake_success)
    lo_dir = tmp_path / "excel-lo"
    lo_dir.mkdir()
    lo_result = excel.run([sample_pdf], {}, lo_dir, reporter=lambda *_args: None)
    assert lo_result.meta["engine"] == "libreoffice"

    monkeypatch.setattr(pdf_to_office, "run_libreoffice_conversion_to_output", lambda *_args, **_kwargs: (False, "lo failed"))
    fallback_dir = tmp_path / "ppt-lo-fallback"
    fallback_dir.mkdir()
    fallback = ppt.run([sample_pdf], {}, fallback_dir)
    assert fallback.meta["fallback_reason"] == "lo failed"


def test_split_modes_validation_groups_and_bookmark_helpers(
    tmp_path: Path,
    sample_pdf: Path,
):
    tool = SplitTool()
    with pytest.raises(ToolError) as invalid:
        tool.validate({"mode": "bad"})
    assert invalid.value.code == ErrorCode.INVALID_PARAMS
    assert split._parse_page_groups("1| 2-3 | ") == ["1", "2-3"]
    assert split._slugify(" Hello 世界.pdf ") == "hello_.pdf"

    range_dir = tmp_path / "split-range"
    range_dir.mkdir()
    grouped = tool.run([sample_pdf], {"mode": "range", "page_groups": "1|2-3"}, range_dir)
    assert grouped.meta["output_count"] == 2

    single_dir = tmp_path / "split-single"
    single_dir.mkdir()
    single = tool.run([sample_pdf], {"mode": "range", "page_range": "1-2"}, single_dir)
    assert single.meta["output_count"] == 1

    each_dir = tmp_path / "split-each"
    each_dir.mkdir()
    each = tool.run([sample_pdf], {"mode": "each_page"}, each_dir, reporter=lambda *_args: None)
    assert each.meta["output_count"] == 5

    chunk_dir = tmp_path / "split-chunk"
    chunk_dir.mkdir()
    chunk = tool.run([sample_pdf], {"mode": "chunk", "chunk_size": 2}, chunk_dir)
    assert chunk.meta["output_count"] == 3

    bookmark_dir = tmp_path / "split-bookmark"
    bookmark_dir.mkdir()
    with pytest.raises(ToolError) as no_bookmarks:
        tool.run([sample_pdf], {"mode": "bookmark"}, bookmark_dir)
    assert no_bookmarks.value.code == ErrorCode.INVALID_PARAMS

    item = SimpleNamespace(destination=2, title="", children=[])
    points: list[tuple[str, int]] = []
    split._walk_outline([item], {}, points)
    assert points == [("bookmark_3", 2)]


def test_signature_visible_and_digital_contracts(
    tmp_path: Path,
    sample_pdf: Path,
    sample_images: list[Path],
    monkeypatch: pytest.MonkeyPatch,
):
    tool = SignatureTool()
    assert tool.validate({"page": 0, "width_pt": 999, "opacity": 2})["opacity"] == 1.0
    with pytest.raises(ToolError) as no_pdf:
        tool.run([sample_images[0]], {}, tmp_path / "sig-no-pdf")
    assert no_pdf.value.code == ErrorCode.INVALID_PARAMS
    with pytest.raises(ToolError) as no_image:
        tool.run([sample_pdf], {}, tmp_path / "sig-no-image")
    assert no_image.value.code == ErrorCode.INVALID_PARAMS
    with pytest.raises(ToolError) as digital_no_cert:
        tool.run([sample_pdf], {"mode": "digital"}, tmp_path / "sig-no-cert")
    assert digital_no_cert.value.code == ErrorCode.INVALID_PARAMS

    visible_dir = tmp_path / "sig-visible"
    visible_dir.mkdir()
    visible = tool.run(
        [sample_pdf, sample_images[0]],
        {"position": "center", "page": 99, "width_pt": 80},
        visible_dir,
    )
    assert visible.meta["mode"] == "visible"
    assert visible.output_files[0].exists()

    cert = tmp_path / "cert.p12"
    cert.write_bytes(b"fake")

    def fake_digital(self, input_pdf: Path, cert_path: Path, output_path: Path, params: dict) -> None:
        assert cert_path == cert
        output_path.write_bytes(input_pdf.read_bytes())

    monkeypatch.setattr(signature.SignatureTool, "_apply_digital_signature", fake_digital)
    digital_dir = tmp_path / "sig-digital"
    digital_dir.mkdir()
    digital = tool.run([sample_pdf, sample_images[0], cert], {"mode": "digital"}, digital_dir)
    assert digital.meta["mode"] == "digital"
    assert digital.output_files[0].exists()


def test_signature_info_detects_fields_and_handles_verification_fallbacks(
    sample_pdf: Path,
    signature_field_pdf: Path,
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
):
    assert SignatureInfoTool().validate({}) == {}
    assert signature_info._verify_signatures(sample_pdf) == []
    monkeypatch.setattr(signature_info, "_verify_signatures", lambda _path: [{"field_name": "Signature1"}])
    result = SignatureInfoTool().run([signature_field_pdf], {}, tmp_path)
    assert result.meta["has_signatures"] is True
    assert result.meta["signature_count"] == 1
    assert result.meta["verification"] == [{"field_name": "Signature1"}]


def test_tile_pages_render_validate_and_run_paths(
    tmp_path: Path,
    sample_pdf: Path,
    monkeypatch: pytest.MonkeyPatch,
):
    tool = TilePagesTool()
    with pytest.raises(ToolError) as invalid:
        tool.validate({"direction": "diagonal"})
    assert invalid.value.code == ErrorCode.INVALID_PARAMS

    monkeypatch.setattr(tile_pages.shutil, "which", lambda name: None)
    assert tile_pages._render_first_page_png(sample_pdf, tmp_path) is None
    with pytest.raises(ToolError) as no_engine:
        tool.run([sample_pdf, sample_pdf], {}, tmp_path / "tile-missing")
    assert no_engine.value.code == ErrorCode.ENGINE_NOT_INSTALLED

    monkeypatch.setattr(tile_pages.shutil, "which", lambda name: "/usr/bin/pdftoppm")
    monkeypatch.setattr(
        tile_pages,
        "run_command",
        lambda *_args, **_kwargs: SimpleNamespace(returncode=1, stderr=b"bad"),
    )
    assert tile_pages._render_first_page_png(sample_pdf, tmp_path) is None

    def fake_render(pdf_path: Path, tmpdir: Path, dpi: int = 96):
        path = tmpdir / f"{pdf_path.stem}.png"
        Image.new("RGB", (20, 20), "white").save(path)
        return path

    monkeypatch.setattr(tile_pages, "_render_first_page_png", fake_render)
    horizontal_dir = tmp_path / "tile-horizontal"
    horizontal_dir.mkdir()
    horizontal = tool.run([sample_pdf, sample_pdf], {"direction": "horizontal"}, horizontal_dir)
    assert horizontal.meta["direction"] == "horizontal"
    assert horizontal.output_files[0].exists()

    vertical_dir = tmp_path / "tile-vertical"
    vertical_dir.mkdir()
    vertical = tool.run(
        [sample_pdf, sample_pdf],
        {"direction": "vertical"},
        vertical_dir,
        reporter=lambda *_args: None,
    )
    assert vertical.meta["direction"] == "vertical"

    monkeypatch.setattr(tile_pages, "_render_first_page_png", lambda *_args, **_kwargs: None)
    fail_dir = tmp_path / "tile-fail"
    fail_dir.mkdir()
    with pytest.raises(ToolError) as render_fail:
        tool.run([sample_pdf, sample_pdf], {}, fail_dir)
    assert render_fail.value.code == ErrorCode.ENGINE_EXEC_FAILED


def test_watermark_image_validation_overlay_positions_and_run_paths(
    tmp_path: Path,
    sample_pdf: Path,
    sample_images: list[Path],
):
    tool = WatermarkImageTool()
    assert tool.validate({"opacity": "0.5", "scale": "0.25", "position": "top_left"})["position"] == "top_left"
    with pytest.raises(ToolError) as missing_inputs:
        tool.run([sample_pdf], {}, tmp_path / "wm-missing")
    assert missing_inputs.value.code == ErrorCode.INVALID_INPUT_FILE

    bad_image = tmp_path / "bad.png"
    bad_image.write_text("bad", encoding="utf-8")
    with pytest.raises(ToolError) as bad:
        tool.run([sample_pdf, bad_image], {}, tmp_path / "wm-bad")
    assert bad.value.code == ErrorCode.INVALID_INPUT_FILE

    for position in ("center", "top_left", "top_right", "bottom_left", "bottom_right"):
        overlay = watermark_image._make_image_overlay(
            Image.new("RGBA", (20, 10), (255, 0, 0, 128)),
            page_w=200,
            page_h=200,
            opacity=0.4,
            scale=0.2,
            position=position,
        )
        assert overlay.getbuffer().nbytes > 0

    paletted = tmp_path / "palette.gif"
    Image.new("P", (30, 30)).save(paletted)
    workdir = tmp_path / "wm-run"
    workdir.mkdir()
    updates: list[int] = []
    result = tool.run(
        [sample_pdf, paletted],
        {"page_range": "1-2", "position": "bottom_right"},
        workdir,
        reporter=lambda percent, message="": updates.append(percent),
    )
    assert result.meta["watermarked_pages"] == 2
    assert result.output_files[0].exists()
    assert updates


def test_annotation_page_tools_and_import_failure_edges(
    tmp_path: Path,
    sample_pdf: Path,
    sample_images: list[Path],
    monkeypatch: pytest.MonkeyPatch,
):
    with pytest.raises(ToolError) as no_blank_targets:
        monkeypatch.setattr(add_blank_pages, "parse_page_range", lambda *_args: [])
        AddBlankPagesTool().run([sample_pdf], {}, tmp_path / "blank-none")
    assert no_blank_targets.value.code == ErrorCode.INVALID_PARAMS
    monkeypatch.setattr(add_blank_pages, "parse_page_range", lambda *_args: [0])
    updates: list[tuple[int, str]] = []
    blank_dir = tmp_path / "blank-reporter"
    blank_dir.mkdir()
    AddBlankPagesTool().run(
        [sample_pdf],
        {"count": 99},
        blank_dir,
        reporter=lambda percent, message="": updates.append((percent, message)),
    )
    assert updates[-1][0] == 100

    assert _make_number_overlay("1", 200, 200, "top_left", 10).getbuffer().nbytes > 0
    assert _make_number_overlay("1", 200, 200, "bottom_right", 10).getbuffer().nbytes > 0

    page_border_tool = PageBorderTool()
    with pytest.raises(ToolError):
        page_border_tool.validate({"border_color": "not-hex"})
    with pytest.raises(ToolError):
        page_border_tool.validate({"bg_color": "nope"})
    bordered_dir = tmp_path / "border-bg"
    bordered_dir.mkdir()
    bordered = page_border_tool.run(
        [sample_pdf],
        {"border_color": "0F0", "bg_color": "#FFEEAA", "page_range": "1"},
        bordered_dir,
    )
    assert bordered.output_files[0].exists()

    stamp_tool = StampTool()
    for position in ("center", "top-left", "top-right", "bottom-left"):
        stamp_dir = tmp_path / f"stamp-{position}"
        stamp_dir.mkdir()
        stamped = stamp_tool.run(
            [sample_pdf, sample_images[0]],
            {"position": position, "page_range": "1"},
            stamp_dir,
        )
        assert stamped.meta["position"] == position

    with pytest.raises(ToolError):
        BarcodeTool().validate({})
    monkeypatch.setitem(sys.modules, "barcode", None)
    with pytest.raises(ToolError) as no_barcode:
        BarcodeTool().run([sample_pdf], {"content": "123"}, tmp_path / "barcode-missing")
    assert no_barcode.value.code == ErrorCode.ENGINE_NOT_INSTALLED

    barcode_module = ModuleType("barcode")
    barcode_module.get_barcode_class = lambda _name: (_ for _ in ()).throw(RuntimeError("bad barcode"))
    barcode_writer = ModuleType("barcode.writer")
    barcode_writer.ImageWriter = object
    monkeypatch.setitem(sys.modules, "barcode", barcode_module)
    monkeypatch.setitem(sys.modules, "barcode.writer", barcode_writer)
    with pytest.raises(ToolError) as barcode_failed:
        BarcodeTool().run([sample_pdf], {"content": "123"}, tmp_path / "barcode-bad")
    assert barcode_failed.value.code == ErrorCode.ENGINE_EXEC_FAILED

    with pytest.raises(ToolError):
        QrCodeTool().validate({})
    monkeypatch.setitem(sys.modules, "qrcode", None)
    with pytest.raises(ToolError) as no_qrcode:
        QrCodeTool().run([sample_pdf], {"content": "hello"}, tmp_path / "qr-missing")
    assert no_qrcode.value.code == ErrorCode.ENGINE_NOT_INSTALLED


def test_external_engine_tool_success_failure_and_validation_edges(
    tmp_path: Path,
    sample_pdf: Path,
    monkeypatch: pytest.MonkeyPatch,
):
    def fake_gs_command(cmd, **_kwargs):
        output_arg = next(part for part in cmd if str(part).startswith("-sOutputFile="))
        Path(str(output_arg).split("=", 1)[1]).write_bytes(b"%PDF-1.4\n%%EOF\n")
        return SimpleNamespace(returncode=0, stdout=b"", stderr=b"")

    with pytest.raises(ToolError):
        CompressTool().validate({"level": "maximum"})
    monkeypatch.setattr(compress.shutil, "which", lambda name: None)
    with pytest.raises(ToolError) as no_gs:
        CompressTool().run([sample_pdf], {}, tmp_path / "compress-missing")
    assert no_gs.value.code == ErrorCode.ENGINE_NOT_INSTALLED
    monkeypatch.setattr(compress.shutil, "which", lambda name: "/usr/bin/gs")
    monkeypatch.setattr(compress, "run_command", fake_gs_command)
    compressed_dir = tmp_path / "compress-ok"
    compressed_dir.mkdir()
    assert CompressTool().run([sample_pdf], {"level": "high"}, compressed_dir).meta["compressed_size"] > 0

    monkeypatch.setattr(linearize.shutil, "which", lambda name: None)
    with pytest.raises(ToolError) as no_qpdf:
        LinearizeTool().run([sample_pdf], {}, tmp_path / "linearize-missing")
    assert no_qpdf.value.code == ErrorCode.ENGINE_NOT_INSTALLED
    monkeypatch.setattr(linearize.shutil, "which", lambda name: "/usr/bin/qpdf")
    monkeypatch.setattr(linearize, "run_command", lambda cmd, **_kwargs: Path(cmd[-1]).write_bytes(sample_pdf.read_bytes()))
    linear_dir = tmp_path / "linearize-ok"
    linear_dir.mkdir()
    assert LinearizeTool().run([sample_pdf], {}, linear_dir, reporter=lambda *_args: None).output_files[0].exists()

    monkeypatch.setattr(repair.shutil, "which", lambda name: None)
    with pytest.raises(ToolError) as no_repair_gs:
        RepairTool().run([sample_pdf], {}, tmp_path / "repair-missing")
    assert no_repair_gs.value.code == ErrorCode.ENGINE_NOT_INSTALLED
    monkeypatch.setattr(repair.shutil, "which", lambda name: "/usr/bin/gs")
    monkeypatch.setattr(repair, "run_command", fake_gs_command)
    repair_dir = tmp_path / "repair-ok"
    repair_dir.mkdir()
    assert RepairTool().run([sample_pdf], {}, repair_dir, reporter=lambda *_args: None).output_files[0].exists()

    pdfa_tool = PdfATool()
    with pytest.raises(ToolError):
        pdfa_tool.validate({"level": "4z"})
    monkeypatch.setattr(pdf_to_pdfa.shutil, "which", lambda name: None)
    with pytest.raises(ToolError) as no_pdfa_gs:
        pdfa_tool.run([sample_pdf], {}, tmp_path / "pdfa-missing")
    assert no_pdfa_gs.value.code == ErrorCode.ENGINE_NOT_INSTALLED
    monkeypatch.setattr(pdf_to_pdfa.shutil, "which", lambda name: "/usr/bin/gs")
    monkeypatch.setattr(pdf_to_pdfa, "run_command", fake_gs_command)
    pdfa_dir = tmp_path / "pdfa-ok"
    pdfa_dir.mkdir()
    assert pdfa_tool.run([sample_pdf], {"level": "3b"}, pdfa_dir, reporter=lambda *_args: None).meta["level"] == "PDF/A-3b"

    validate_tool = ValidateTool()
    monkeypatch.setattr(validate.shutil, "which", lambda name: None)
    with pytest.raises(ToolError):
        validate_tool.run([sample_pdf], {}, tmp_path)
    monkeypatch.setattr(validate.shutil, "which", lambda name: "/usr/bin/qpdf")

    def fake_qpdf(cmd, **_kwargs):
        if "--check-linearization" in cmd:
            raise RuntimeError("linearization check failed")
        return SimpleNamespace(returncode=1, stdout=b"PDF Version: 1.7\nwarning: bad xref\n", stderr=b"")

    monkeypatch.setattr(validate, "run_command", fake_qpdf)
    validate_dir = tmp_path / "validate"
    validate_dir.mkdir()
    validation = validate_tool.run([sample_pdf], {}, validate_dir)
    assert validation.meta["is_valid"] is False
    assert validation.meta["is_linearized"] is False
    assert validation.meta["issues"] == ["warning: bad xref"]


def test_compare_text_image_diff_and_missing_page_branches(
    tmp_path: Path,
    sample_pdf: Path,
    monkeypatch: pytest.MonkeyPatch,
):
    one_page = tmp_path / "one-page.pdf"
    with pikepdf.open(sample_pdf) as src:
        out = pikepdf.Pdf.new()
        out.pages.append(src.pages[0])
        out.save(one_page)

    def fake_page_image(path: Path, page_idx: int, dpi: int = 72):
        if path == one_page and page_idx > 0:
            return None
        size = (10, 10) if path == one_page else (20, 20)
        color = "black" if path == one_page else "white"
        return Image.new("RGB", size, color)

    monkeypatch.setattr(compare, "_render_page_png", fake_page_image)
    monkeypatch.setattr(compare, "_extract_page_text", lambda page: f"text-{id(page)}")
    workdir = tmp_path / "compare-diff"
    workdir.mkdir()
    result = CompareTool().run([sample_pdf, one_page], {"highlight_color": "yellow"}, workdir)
    assert result.meta["pages_with_diff"] == 5
    assert result.meta["pages_with_text_diff"] >= 1


def test_signature_digital_and_info_optional_dependency_edges(
    tmp_path: Path,
    sample_pdf: Path,
    signature_field_pdf: Path,
    monkeypatch: pytest.MonkeyPatch,
):
    cert = tmp_path / "cert.p12"
    cert.write_bytes(b"fake-cert")

    real_import = builtins.__import__

    def reject_pyhanko_import(name, *args, **kwargs):
        if name.startswith("pyhanko"):
            raise ImportError("pyhanko missing")
        return real_import(name, *args, **kwargs)

    monkeypatch.setattr(builtins, "__import__", reject_pyhanko_import)
    with pytest.raises(ToolError) as no_pyhanko:
        SignatureTool()._apply_digital_signature(sample_pdf, cert, tmp_path / "signed-missing.pdf", SignatureTool().validate({}))
    assert no_pyhanko.value.code == ErrorCode.ENGINE_NOT_INSTALLED
    monkeypatch.setattr(builtins, "__import__", real_import)

    pyhanko = ModuleType("pyhanko")
    pdf_utils = ModuleType("pyhanko.pdf_utils")
    incremental_writer = ModuleType("pyhanko.pdf_utils.incremental_writer")
    sign_module = ModuleType("pyhanko.sign")
    signers = ModuleType("pyhanko.sign.signers")

    class _Writer:
        def __init__(self, stream) -> None:
            self.stream = stream

    class _SimpleSigner:
        @staticmethod
        def load_pkcs12(path: Path, passphrase: bytes | None = None):
            assert path == cert
            assert passphrase == b"secret"
            return "signer"

    class _PdfSigner:
        def __init__(self, meta, signer) -> None:
            self.meta = meta
            self.signer = signer

        def sign_pdf(self, writer, *, output) -> None:
            assert isinstance(writer, _Writer)
            output.write(b"signed")

    incremental_writer.IncrementalPdfFileWriter = _Writer
    signers.SimpleSigner = _SimpleSigner
    signers.PdfSignatureMetadata = lambda **kwargs: SimpleNamespace(**kwargs)
    signers.PdfSigner = _PdfSigner
    sign_module.signers = signers
    monkeypatch.setitem(sys.modules, "pyhanko", pyhanko)
    monkeypatch.setitem(sys.modules, "pyhanko.pdf_utils", pdf_utils)
    monkeypatch.setitem(sys.modules, "pyhanko.pdf_utils.incremental_writer", incremental_writer)
    monkeypatch.setitem(sys.modules, "pyhanko.sign", sign_module)
    monkeypatch.setitem(sys.modules, "pyhanko.sign.signers", signers)

    signed = tmp_path / "signed.pdf"
    SignatureTool()._apply_digital_signature(
        sample_pdf,
        cert,
        signed,
        SignatureTool().validate({"p12_password": "secret", "reason": "ok", "location": "lab"}),
    )
    assert signed.read_bytes() == b"signed"

    monkeypatch.setitem(sys.modules, "pyhanko", None)
    assert signature_info._verify_signatures(sample_pdf) == []

    class _EmbeddedSig:
        field_name = "Signature1"

    class _Reader:
        def __init__(self, _fh) -> None:
            self.embedded_signatures = [_EmbeddedSig()]

    validation = ModuleType("pyhanko.sign.validation")
    validation.ValidationContext = lambda allow_fetching=False: SimpleNamespace(allow_fetching=allow_fetching)
    validation.validate_pdf_signature = lambda _sig, _ctx: SimpleNamespace(
        intact=True,
        trusted=False,
        bottom_line="signature valid",
    )
    reader_module = ModuleType("pyhanko.pdf_utils.reader")
    reader_module.PdfFileReader = _Reader
    monkeypatch.setitem(sys.modules, "pyhanko", pyhanko)
    monkeypatch.setitem(sys.modules, "pyhanko.pdf_utils.reader", reader_module)
    monkeypatch.setitem(sys.modules, "pyhanko.sign.validation", validation)
    verified = signature_info._verify_signatures(sample_pdf)
    assert verified == [
        {
            "field_name": "Signature1",
            "intact": True,
            "trusted": False,
            "bottom_line": "signature valid",
        }
    ]

    assert SignatureInfoTool().run([signature_field_pdf], {}, tmp_path).meta["signature_count"] == 1


def test_page_operation_validation_and_reporter_edges(
    tmp_path: Path,
    sample_pdf: Path,
    encrypted_pdf: Path,
):
    with pytest.raises(ToolError):
        DeleteTool().validate({})
    with pytest.raises(ToolError):
        CropTool().validate({"top": -1})
    with pytest.raises(ToolError):
        ResizeTool().validate({"target_size": "Poster"})
    with pytest.raises(ToolError):
        ReorderTool().validate({})
    with pytest.raises(ToolError):
        ReorderTool().validate({"order": object()})
    with pytest.raises(ToolError):
        ReorderTool().validate({"order": []})
    with pytest.raises(ToolError):
        ReorderTool().run([sample_pdf], {"order": "6"}, tmp_path / "reorder-oob")

    reorder_dir = tmp_path / "reorder-list"
    reorder_dir.mkdir()
    assert ReorderTool().run([sample_pdf], {"order": [1, 3, 5]}, reorder_dir).meta["output_pages"] == 3

    with pytest.raises(ToolError):
        DecryptTool().validate({})
    with pytest.raises(ToolError) as wrong_password:
        DecryptTool().run([encrypted_pdf], {"password": "bad"}, tmp_path / "decrypt-bad")
    assert wrong_password.value.code == ErrorCode.INVALID_PARAMS

    with pytest.raises(ToolError):
        EncryptTool().validate({})

    crop_dir = tmp_path / "crop-reporter"
    crop_dir.mkdir()
    crop_updates: list[int] = []
    CropTool().run([sample_pdf], {"left": 1, "page_range": "1"}, crop_dir, reporter=lambda percent: crop_updates.append(percent))
    assert crop_updates == [20]
    with pytest.raises(ToolError):
        CropTool().run([sample_pdf], {"left": 9999, "page_range": "1"}, tmp_path / "crop-too-large")

    resize_dir = tmp_path / "resize-reporter"
    resize_dir.mkdir()
    resize_updates: list[int] = []
    ResizeTool().run([sample_pdf], {"target_size": "A4", "page_range": "1"}, resize_dir, reporter=lambda percent: resize_updates.append(percent))
    assert resize_updates == [20]

    with pytest.raises(ToolError):
        WatermarkTextTool().validate({})
    watermark_dir = tmp_path / "watermark-text"
    watermark_dir.mkdir()
    watermark_updates: list[int] = []
    watermark = WatermarkTextTool().run(
        [sample_pdf],
        {"text": "TEST", "color": "unknown", "page_range": "1"},
        watermark_dir,
        reporter=lambda percent: watermark_updates.append(percent),
    )
    assert watermark.output_files[0].exists()
    assert watermark_updates == [20]

    with pytest.raises(ToolError):
        HeaderFooterTool().validate({})
    header_dir = tmp_path / "header-footer"
    header_dir.mkdir()
    assert HeaderFooterTool().run(
        [sample_pdf],
        {"header": "H {page}/{total}", "footer": "F", "page_range": "1"},
        header_dir,
    ).output_files[0].exists()


def test_merge_nup_pages_zip_and_text_conversion_edges(
    tmp_path: Path,
    sample_pdf: Path,
    monkeypatch: pytest.MonkeyPatch,
):
    with pytest.raises(ToolError):
        MergeTool().validate({"mode": "shuffle"})
    with pytest.raises(ToolError):
        MergeTool().run([sample_pdf], {}, tmp_path / "merge-one")
    merge_updates: list[tuple[int, str]] = []
    merge_dir = tmp_path / "merge-reporter"
    merge_dir.mkdir()
    MergeTool().run(
        [sample_pdf, sample_pdf],
        {},
        merge_dir,
        reporter=lambda percent, message="": merge_updates.append((percent, message)),
    )
    assert merge_updates[-1][0] == 100
    insert_dir = tmp_path / "merge-insert"
    insert_dir.mkdir()
    MergeTool().run([sample_pdf, sample_pdf], {"mode": "insert", "insert_position": 99}, insert_dir)

    nup_tool = NUpTool()
    with pytest.raises(ToolError):
        nup_tool.validate({"layout": "3-up"})
    monkeypatch.setattr(nup.shutil, "which", lambda name: None)
    assert nup._render_page_to_png(sample_pdf, 0, tmp_path) is None
    with pytest.raises(ToolError):
        nup_tool.run([sample_pdf], {}, tmp_path / "nup-missing")
    monkeypatch.setattr(nup.shutil, "which", lambda name: "/usr/bin/pdftoppm")
    monkeypatch.setattr(nup, "run_command", lambda *_args, **_kwargs: SimpleNamespace(returncode=1))
    assert nup._render_page_to_png(sample_pdf, 0, tmp_path) is None
    monkeypatch.setattr(nup, "run_command", lambda *_args, **_kwargs: SimpleNamespace(returncode=0))
    assert nup._render_page_to_png(sample_pdf, 1, tmp_path) is None

    def fake_nup_render(_pdf: Path, page_idx: int, tmpdir: Path, dpi: int = 96):
        if page_idx == 0:
            return None
        path = tmpdir / f"nup-{page_idx}.png"
        Image.new("RGB", (30, 20), "white").save(path)
        return path

    monkeypatch.setattr(nup, "_render_page_to_png", fake_nup_render)
    nup_dir = tmp_path / "nup-ok"
    nup_dir.mkdir()
    nup_updates: list[int] = []
    nup_result = nup_tool.run(
        [sample_pdf],
        {"layout": "4-up", "orientation": "landscape"},
        nup_dir,
        reporter=lambda percent, message="": nup_updates.append(percent),
    )
    assert nup_result.meta["output_sheets"] == 2
    assert nup_updates[-1] == 100

    zip_tool = PagesToZipTool()
    with pytest.raises(ToolError):
        zip_tool.validate({"format": "bmp"})
    monkeypatch.setattr(pages_to_zip.shutil, "which", lambda name: None)
    with pytest.raises(ToolError):
        zip_tool.run([sample_pdf], {}, tmp_path / "zip-missing")
    monkeypatch.setattr(pages_to_zip.shutil, "which", lambda name: "/usr/bin/pdftoppm")

    def fake_zip_render(cmd, **_kwargs):
        Path(cmd[-1]).parent.joinpath("page-1.unexpected").write_bytes(b"img")

    monkeypatch.setattr(pages_to_zip, "run_command", fake_zip_render)
    zip_dir = tmp_path / "zip-ok"
    zip_dir.mkdir()
    zip_updates: list[tuple[int, str]] = []
    zipped = zip_tool.run([sample_pdf], {"format": "png", "dpi": 999}, zip_dir, reporter=lambda p, m="": zip_updates.append((p, m)))
    assert zipped.meta["page_count"] == 1
    assert zip_updates[-1] == (100, "Done")

    text_updates: list[int] = []
    text_dir = tmp_path / "pdf-text"
    text_dir.mkdir()
    text_result = PdfToTextTool().run([sample_pdf], {"page_range": "1"}, text_dir, reporter=lambda percent: text_updates.append(percent))
    assert text_result.meta["pages_extracted"] == 1
    assert text_updates == [100]

    monkeypatch.setattr(pdf_to_text.pikepdf, "parse_content_stream", lambda _page: [([pikepdf.Array([pikepdf.String("A")])], "TJ")])
    with pikepdf.open(sample_pdf) as pdf:
        assert pdf_to_text._extract_page_text(pdf.pages[0]) == "A"
    monkeypatch.setattr(pdf_to_text.pikepdf, "parse_content_stream", lambda _page: (_ for _ in ()).throw(RuntimeError("bad stream")))
    with pikepdf.open(sample_pdf) as pdf, pytest.raises(ToolError):
        pdf_to_text._extract_page_text(pdf.pages[0])


def test_markdown_word_office_and_metadata_edges(
    tmp_path: Path,
    sample_pdf: Path,
    sample_docx: Path,
    sample_xlsx: Path,
    sample_pptx: Path,
    monkeypatch: pytest.MonkeyPatch,
):
    assert PdfToWordTool().validate({}) == {}
    monkeypatch.setattr(pdf_to_word.shutil, "which", lambda name: "/usr/bin/libreoffice")
    monkeypatch.setattr(
        pdf_to_word,
        "run_libreoffice_conversion_to_output",
        lambda *_args, **_kwargs: (False, "lo failed"),
    )
    monkeypatch.setattr(pdf_to_word, "_extract_page_text", lambda _page: "")
    word_dir = tmp_path / "word-fallback"
    word_dir.mkdir()
    word_updates: list[tuple[int, str]] = []
    word = PdfToWordTool().run([sample_pdf], {}, word_dir, reporter=lambda p, m="": word_updates.append((p, m)))
    assert word.meta["fallback_reason"] == "lo failed"
    assert word_updates[-1] == (100, "Conversion complete")

    real_import = builtins.__import__

    def reject_docx_import(name, *args, **kwargs):
        if name == "docx":
            raise ImportError("docx missing")
        return real_import(name, *args, **kwargs)

    monkeypatch.setattr(builtins, "__import__", reject_docx_import)
    with pytest.raises(ToolError):
        PdfToWordTool._fallback_convert(sample_pdf, tmp_path / "missing-docx.docx")
    monkeypatch.setattr(builtins, "__import__", real_import)

    fake_pdfminer_high = ModuleType("pdfminer.high_level")
    fake_pdfminer_layout = ModuleType("pdfminer.layout")

    class _FakeText:
        def __init__(self, text: str) -> None:
            self._text = text

        def get_text(self) -> str:
            return self._text

    fake_pdfminer_high.extract_pages = lambda _path: [[_FakeText("a\nb")], [_FakeText("c")]]
    fake_pdfminer_layout.LTTextContainer = _FakeText
    monkeypatch.setitem(sys.modules, "pdfminer.high_level", fake_pdfminer_high)
    monkeypatch.setitem(sys.modules, "pdfminer.layout", fake_pdfminer_layout)
    markdown_dir = tmp_path / "markdown"
    markdown_dir.mkdir()
    md_updates: list[tuple[int, str]] = []
    md = PdfToMarkdownTool().run(
        [sample_pdf],
        {"preserve_layout": True},
        markdown_dir,
        reporter=lambda p, m="": md_updates.append((p, m)),
    )
    assert md.meta["pages"] == 2
    assert md_updates[-1] == (100, "Done")

    fake_pdfminer_high.extract_pages = lambda _path: (_ for _ in ()).throw(RuntimeError("extract failed"))
    with pytest.raises(ToolError) as md_failed:
        PdfToMarkdownTool().run([sample_pdf], {}, tmp_path / "markdown-fail")
    assert md_failed.value.code == ErrorCode.ENGINE_EXEC_FAILED

    office_tool = OfficeToPdfTool()
    with pytest.raises(ToolError):
        office_tool._fallback_convert(tmp_path / "input.txt", tmp_path / "out.pdf")
    office_xlsx = tmp_path / "office-xlsx.pdf"
    office_tool._fallback_convert(sample_xlsx, office_xlsx)
    assert office_xlsx.exists()
    office_pptx = tmp_path / "office-pptx.pdf"
    office_tool._fallback_convert(sample_pptx, office_pptx)
    assert office_pptx.exists()

    from docx import Document

    table_doc = tmp_path / "table.docx"
    doc = Document()
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "A"
    table.cell(0, 1).text = "B"
    doc.save(table_doc)
    table_pdf = tmp_path / "table.pdf"
    office_tool._fallback_convert(table_doc, table_pdf)
    assert table_pdf.exists()

    monkeypatch.setattr(office_to_pdf.shutil, "which", lambda name: None)
    monkeypatch.setattr(OfficeToPdfTool, "_fallback_convert", staticmethod(lambda _inp, _out: None))
    with pytest.raises(ToolError):
        office_tool.run([sample_docx], {}, tmp_path / "office-no-output")

    meta_pdf = tmp_path / "meta.pdf"
    with pikepdf.open(sample_pdf) as pdf:
        pdf.docinfo["/Author"] = "Tester"
        pdf.Root.Metadata = pdf.make_stream(b"<xmp></xmp>")
        pdf.trailer["/CreationDate"] = "D:20260101000000"
        pdf.save(meta_pdf)
    removed_dir = tmp_path / "remove-metadata"
    removed_dir.mkdir()
    removed = RemoveMetadataTool().run([meta_pdf], {}, removed_dir)
    assert "/Author" in removed.meta["removed_fields"]
    assert RemoveMetadataTool().validate({}) == {}
    assert MetadataInfoTool().validate({}) == {}
    assert MetadataInfoTool().run([meta_pdf], {}, tmp_path).meta["docinfo"]["/Author"] == "Tester"


def test_blank_pages_and_embedded_resource_edges(
    tmp_path: Path,
    sample_pdf: Path,
    attachment_pdf: Path,
    image_pdf: Path,
    monkeypatch: pytest.MonkeyPatch,
):
    assert RemoveBlankPagesTool().validate({"threshold": 2})["threshold"] == 1.0
    import pdf_agent.external_commands as external_commands

    monkeypatch.setattr(shutil, "which", lambda name: None)
    with pikepdf.open(sample_pdf) as pdf:
        assert remove_blank_pages._is_blank(pdf.pages[0]) is False

    monkeypatch.setattr(shutil, "which", lambda name: "/usr/bin/gs")
    monkeypatch.setattr(external_commands, "run_command", lambda *_args, **_kwargs: SimpleNamespace(returncode=1))
    with pikepdf.open(sample_pdf) as pdf:
        assert remove_blank_pages._is_blank(pdf.pages[0]) is False

    def fake_blank_command(cmd, **_kwargs):
        output = next(part for part in cmd if str(part).startswith("-sOutputFile="))
        Image.new("L", (2, 2), 255).save(Path(str(output).split("=", 1)[1]))
        return SimpleNamespace(returncode=0)

    monkeypatch.setattr(external_commands, "run_command", fake_blank_command)
    with pikepdf.open(sample_pdf) as pdf:
        assert remove_blank_pages._is_blank(pdf.pages[0], threshold=0.9) is True

    monkeypatch.setattr(remove_blank_pages, "_is_blank", lambda *_args, **_kwargs: True)
    with pytest.raises(ToolError):
        RemoveBlankPagesTool().run([sample_pdf], {}, tmp_path / "blank-all")

    monkeypatch.setattr(remove_blank_pages, "_is_blank", lambda *_args, **_kwargs: False)
    no_blank_dir = tmp_path / "blank-none-run"
    no_blank_dir.mkdir()
    blank_updates: list[tuple[int, str]] = []
    result = RemoveBlankPagesTool().run([sample_pdf], {}, no_blank_dir, reporter=lambda p, m="": blank_updates.append((p, m)))
    assert result.meta["removed_pages"] == 0
    assert blank_updates[-1] == (100, "Done")

    assert ExtractAttachmentsTool().validate({}) == {}
    attachment_dir = tmp_path / "attachments"
    attachment_dir.mkdir()
    extracted_attachment = ExtractAttachmentsTool().run([attachment_pdf], {}, attachment_dir, reporter=lambda *_args: None)
    assert extracted_attachment.meta["attachment_count"] == 1

    assert ExtractImagesTool().validate({}) == {}
    image_dir = tmp_path / "images"
    image_dir.mkdir()
    extracted_images = ExtractImagesTool().run([image_pdf], {}, image_dir, reporter=lambda *_args: None)
    assert extracted_images.meta["image_count"] >= 1


def test_auto_rotate_pdf_text_ocr_and_run_error_edges(
    tmp_path: Path,
    sample_pdf: Path,
    sample_images: list[Path],
    monkeypatch: pytest.MonkeyPatch,
):
    real_import = builtins.__import__

    def reject_pdfminer_import(name, *args, **kwargs):
        if name.startswith("pdfminer"):
            raise ImportError("pdfminer missing")
        return real_import(name, *args, **kwargs)

    monkeypatch.setattr(builtins, "__import__", reject_pdfminer_import)
    assert auto_rotate._detect_rotation_from_pdf_text(sample_pdf, 0) == (0, None)
    monkeypatch.setattr(builtins, "__import__", real_import)

    high_level = ModuleType("pdfminer.high_level")
    layout = ModuleType("pdfminer.layout")

    class _Char:
        def __init__(self, text: str, matrix: tuple[int, int, int, int, int, int]) -> None:
            self._text = text
            self.matrix = matrix

        def get_text(self) -> str:
            return self._text

    class _Container(list):
        pass

    layout.LTChar = _Char
    layout.LTContainer = _Container
    high_level.extract_pages = lambda *_args, **_kwargs: [_Container([_Char("abc", (0, 1, 0, 0, 0, 0))])]
    monkeypatch.setitem(sys.modules, "pdfminer.high_level", high_level)
    monkeypatch.setitem(sys.modules, "pdfminer.layout", layout)
    assert auto_rotate._detect_rotation_from_pdf_text(sample_pdf, 0) == (90, 3.0)

    high_level.extract_pages = lambda *_args, **_kwargs: [_Container([_Char("!!!", (0, 1, 0, 0, 0, 0))])]
    assert auto_rotate._detect_rotation_from_pdf_text(sample_pdf, 0) == (0, None)

    high_level.extract_pages = lambda *_args, **_kwargs: (_ for _ in ()).throw(RuntimeError("bad pdfminer"))
    assert auto_rotate._detect_rotation_from_pdf_text(sample_pdf, 0) == (0, None)

    monkeypatch.setattr(auto_rotate.shutil, "which", lambda name: "/usr/bin/tesseract")
    monkeypatch.setattr(
        auto_rotate,
        "run_command",
        lambda *_args, **_kwargs: SimpleNamespace(returncode=1, stdout=b"", stderr=b"bad"),
    )
    assert auto_rotate._ocr_rotation_score(sample_images[0], 0) is None

    monkeypatch.setattr(
        auto_rotate,
        "run_command",
        lambda *_args, **_kwargs: SimpleNamespace(returncode=0, stdout=b"", stderr=b""),
    )
    assert auto_rotate._ocr_rotation_score(sample_images[0], 0) is None

    monkeypatch.setattr(
        auto_rotate,
        "run_command",
        lambda *_args, **_kwargs: SimpleNamespace(returncode=0, stdout=b"level\ttext\n1\ta\n", stderr=b""),
    )
    assert auto_rotate._ocr_rotation_score(sample_images[0], 0) is None

    monkeypatch.setattr(
        auto_rotate,
        "run_command",
        lambda *_args, **_kwargs: SimpleNamespace(returncode=0, stdout=b"level\tconf\ttext\n1\t90\n2\t91\t\n3\t91\tabc\n", stderr=b""),
    )
    assert auto_rotate._ocr_rotation_score(sample_images[0], 0) is None

    monkeypatch.setattr(auto_rotate.shutil, "which", lambda name: f"/usr/bin/{name}")
    monkeypatch.setattr(auto_rotate, "_render_page_png", lambda *_args, **_kwargs: None)
    with pytest.raises(ToolError) as render_missing:
        AutoRotateTool().run([sample_pdf], {}, tmp_path / "auto-render-missing")
    assert render_missing.value.code == ErrorCode.OUTPUT_GENERATION_FAILED

    monkeypatch.setattr(auto_rotate, "_render_page_png", lambda _pdf, index, tmpdir: sample_images[0])
    monkeypatch.setattr(
        auto_rotate,
        "_detect_rotation",
        lambda _image: (_ for _ in ()).throw(ToolError(ErrorCode.ENGINE_EXEC_FAILED, "fatal osd")),
    )
    with pytest.raises(ToolError) as fatal_osd:
        AutoRotateTool().run([sample_pdf], {}, tmp_path / "auto-fatal")
    assert fatal_osd.value.message == "fatal osd"


def test_redact_text_query_failure_nonmatch_and_reporter_edges(
    tmp_path: Path,
    sample_pdf: Path,
    monkeypatch: pytest.MonkeyPatch,
):
    monkeypatch.setattr(redact.shutil, "which", lambda name: "/usr/bin/pdftotext")
    monkeypatch.setattr(
        redact,
        "run_command",
        lambda *_args, **_kwargs: SimpleNamespace(returncode=1, stderr=b"pdftotext bad"),
    )
    with pytest.raises(ToolError) as text_lookup_failed:
        redact._find_text_boxes(sample_pdf, "secret")
    assert text_lookup_failed.value.code == ErrorCode.ENGINE_EXEC_FAILED

    def fake_nonmatching_pdftotext(cmd, **_kwargs):
        Path(cmd[-1]).write_text(
            '<doc><page height="100"><word xMin="1" yMin="2" xMax="3" yMax="4">other</word></page></doc>',
            encoding="utf-8",
        )
        return SimpleNamespace(returncode=0, stderr=b"")

    monkeypatch.setattr(redact, "run_command", fake_nonmatching_pdftotext)
    assert redact._find_text_boxes(sample_pdf, "secret") == []

    monkeypatch.setattr(redact, "_find_text_boxes", lambda *_args: [{"page": 1, "x": 1, "y": 2, "width": 3, "height": 4}])
    monkeypatch.setattr(redact.shutil, "which", lambda name: None)
    redact_dir = tmp_path / "redact-query"
    redact_dir.mkdir()
    updates: list[int] = []
    result = RedactTool().run(
        [sample_pdf],
        {"text_query": "secret", "fill_color": "white"},
        redact_dir,
        reporter=lambda percent, message="": updates.append(percent),
    )
    assert result.meta["redaction_count"] == 1
    assert updates == [10]

    skip_dir = tmp_path / "redact-skip"
    skip_dir.mkdir()
    skipped = RedactTool().run(
        [sample_pdf],
        {"regions_json": '[{"page":1,"x":1,"y":1,"width":2,"height":2}]', "page_range": "2"},
        skip_dir,
    )
    assert skipped.meta["redaction_mode"] == "visual_only"


def test_form_fill_helpers_and_builtin_misc_edges(
    tmp_path: Path,
    sample_pdf: Path,
    monkeypatch: pytest.MonkeyPatch,
):
    class _Field(dict):
        def get(self, key, default=None):
            if key == "/broken":
                raise RuntimeError("broken")
            return super().get(key, default)

    assert form_fill._get_form_fields(SimpleNamespace(Root={pikepdf.Name("/AcroForm"): {}})) == {}
    fake_pdf = SimpleNamespace(
        Root={
            pikepdf.Name("/AcroForm"): {
                pikepdf.Name("/Fields"): [
                    {"/T": "text", "/V": 123},
                    object(),
                ]
            }
        }
    )
    assert form_fill._get_form_fields(fake_pdf)["text"] == "123"

    moved = []
    monkeypatch.setattr(form_fill, "shutil", SimpleNamespace(move=lambda src, dst: moved.append((src, dst))))

    class _Flatten:
        def run(self, _inputs, _params, output_dir):
            generated = output_dir / "generated.pdf"
            generated.write_bytes(b"%PDF-1.4\n%%EOF\n")
            return SimpleNamespace(output_files=[generated])

    import pdf_agent.tools._builtins.flatten as flatten_module

    monkeypatch.setattr(flatten_module, "FlattenTool", _Flatten)
    form_fill._flatten_pdf(sample_pdf, tmp_path / "flattened.pdf")
    assert moved

    with pytest.raises(ToolError):
        FlattenTool().run([sample_pdf], {}, tmp_path / "flatten-missing")
    monkeypatch.setattr(flatten.shutil, "which", lambda name: "/usr/bin/gs")
    monkeypatch.setattr(flatten, "run_command", lambda cmd, **_kwargs: Path(str(cmd[-2]).split("=", 1)[1]).write_bytes(sample_pdf.read_bytes()))
    flatten_dir = tmp_path / "flatten-ok"
    flatten_dir.mkdir()
    assert FlattenTool().run([sample_pdf], {}, flatten_dir).output_files[0].exists()

    booklet_dir = tmp_path / "booklet"
    booklet_dir.mkdir()
    booklet_updates: list[int] = []
    booklet_result = BookletTool().run([sample_pdf], {}, booklet_dir, reporter=lambda percent, message="": booklet_updates.append(percent))
    assert booklet_result.meta["booklet_pages"] % 4 == 0
    assert booklet_updates[-1] == 100

    with pytest.raises(ToolError):
        ExtractTool().validate({})
    extract_dir = tmp_path / "extract"
    extract_dir.mkdir()
    extract_updates: list[int] = []
    assert ExtractTool().run([sample_pdf], {"page_range": "1"}, extract_dir, reporter=lambda percent: extract_updates.append(percent)).meta["extracted_pages"] == 1
    assert extract_updates == [100]


def test_office_conversion_optional_dependency_edges(
    tmp_path: Path,
    sample_pdf: Path,
    sample_docx: Path,
    monkeypatch: pytest.MonkeyPatch,
):
    real_import = builtins.__import__

    def reject_optional_office_deps(name, *args, **kwargs):
        if name in {"docx", "openpyxl", "pptx"} or name.startswith("reportlab"):
            raise ImportError(f"{name} missing")
        return real_import(name, *args, **kwargs)

    monkeypatch.setattr(builtins, "__import__", reject_optional_office_deps)
    with pytest.raises(ToolError):
        office_to_pdf._docx_to_pdf_fallback(sample_docx, tmp_path / "docx.pdf")
    with pytest.raises(ToolError):
        office_to_pdf._xlsx_to_pdf_fallback(tmp_path / "sheet.xlsx", tmp_path / "sheet.pdf")
    with pytest.raises(ToolError):
        office_to_pdf._pptx_to_pdf_fallback(tmp_path / "deck.pptx", tmp_path / "deck.pdf")
    with pytest.raises(ToolError):
        office_to_pdf._render_text_lines_to_pdf(["hello"], tmp_path / "text.pdf")
    with pytest.raises(ToolError):
        PdfToExcelTool._fallback_convert(sample_pdf, tmp_path / "out.xlsx")
    with pytest.raises(ToolError):
        PdfToPptTool._fallback_convert(sample_pdf, tmp_path / "out.pptx")
    monkeypatch.setattr(builtins, "__import__", real_import)

    monkeypatch.setattr(pdf_to_office, "_extract_page_text", lambda _page: "")
    excel_out = tmp_path / "blank.xlsx"
    PdfToExcelTool._fallback_convert(sample_pdf, excel_out)
    assert excel_out.exists()

    monkeypatch.setattr(pdf_to_office, "run_libreoffice_conversion_to_output", lambda *_args, **_kwargs: (True, None))
    monkeypatch.setattr(pdf_to_office.shutil, "which", lambda name: "/usr/bin/libreoffice")
    ppt_dir = tmp_path / "ppt-lo"
    ppt_dir.mkdir()
    ppt = PdfToPptTool().run([sample_pdf], {}, ppt_dir, reporter=lambda *_args: None)
    assert ppt.meta["engine"] == "libreoffice"


def test_signature_info_parse_exceptions_and_import_failure_edges(
    tmp_path: Path,
    sample_pdf: Path,
    monkeypatch: pytest.MonkeyPatch,
):
    real_import = builtins.__import__

    def reject_pyhanko_import(name, *args, **kwargs):
        if name.startswith("pyhanko"):
            raise ImportError("pyhanko missing")
        return real_import(name, *args, **kwargs)

    monkeypatch.setattr(builtins, "__import__", reject_pyhanko_import)
    assert signature_info._verify_signatures(sample_pdf) == []
    monkeypatch.setattr(builtins, "__import__", real_import)

    class _BadStr:
        def __str__(self) -> str:
            raise RuntimeError("cannot stringify")

    class _Field:
        def __init__(self, *, fail: bool = False) -> None:
            self.fail = fail

        def get(self, key, default=None):
            if self.fail:
                raise RuntimeError("bad field")
            if key == "/FT":
                return "/Sig"
            if key == "/T":
                return "SignatureBadValue"
            if key == "/V":
                return {"/Name": _BadStr()}
            return default

    class _AcroForm:
        def get(self, key, default=None):
            if key == "/Fields":
                return [_Field(), _Field(fail=True)]
            return default

    class _Root:
        def __contains__(self, key):
            return key == pikepdf.Name("/AcroForm")

        def __getitem__(self, key):
            assert key == "/AcroForm"
            return _AcroForm()

    class _Pdf:
        Root = _Root()

        def __enter__(self):
            return self

        def __exit__(self, exc_type, exc, tb):
            return False

    monkeypatch.setattr(signature_info.pikepdf, "open", lambda _path: _Pdf())
    monkeypatch.setattr(signature_info, "_verify_signatures", lambda _path: [])
    result = SignatureInfoTool().run([tmp_path / "fake.pdf"], {}, tmp_path)
    assert result.meta["signatures"] == [{"field": "SignatureBadValue"}]
